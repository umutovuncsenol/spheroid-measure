"""Excel export. Only user-ACCEPTED objects go to the 'Measurements' sheet; everything
unresolved (unreviewed, rejected, wells marked not measurable, wells without candidates,
unprocessed / unreviewed / flagged images) goes to 'Review log'.

Values are written as numbers computed by app/measure.py (formulas documented in
'Settings & definitions'). Physical columns stay EMPTY unless the image's calibration is
confirmed; pixel columns are always filled.
"""
import json
import os
import shutil

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw, ImageFont

from . import measure as M
from .store import APP_VERSION, load_config, load_rgb8, now, contours_xy, disk_name, region_kind

REGIONS = [('cells', 'Cell-covered'), ('envelope', 'Envelope'), ('core', 'Dense core')]
AV_NOTE = {(True, 'rejected'): 'rejected candidate: no correction numbers (an unchanged rejected mask would give IoU 1)',
           (True, 'unreviewed'): 'not reviewed yet: no correction numbers',
           (False, 'accepted'): 'added by the reviewer (missed by the automatic step): no automatic mask',
           (False, 'rejected'): 'added by the reviewer, then rejected',
           (False, 'unreviewed'): 'added by the reviewer, not reviewed yet'}
EFFORT_HEAD = ['Original filename', 'Image ID', 'Image review complete',
               'Measured active review time (s)', 'Activity window (s)',
               'Rough estimate from action gaps (s)', 'Idle cap per gap (s)', 'Reviewer events (n)',
               'First event', 'Last event', 'Auto candidates', 'Accepted unchanged', 'Accepted after edit',
               'Rejected candidates', 'Unreviewed objects', 'Added by reviewer: drawn (missed)',
               'Added by reviewer: from detected parts (missed)', 'Added by reviewer and accepted',
               'Mean IoU auto vs final (accepted auto objects)', 'Accepted auto objects with IoU < 0.9',
               'Event counts (JSON)']
REGION_TXT = {'well': 'well', 'grid cell': 'grid cell (working region, not a well)', 'whole image': 'whole image', '': ''}
FONT = Font(name='Arial', size=10)
BOLD = Font(name='Arial', size=10, bold=True)
HEAD_FILL = PatternFill('solid', fgColor='DDE4EE')


def _font(size):
    for p in ('/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf', '/Library/Fonts/Arial Bold.ttf',
              '/System/Library/Fonts/Supplemental/Arial Bold.ttf', 'C:/Windows/Fonts/arialbd.ttf'):
        if os.path.exists(p):
            return ImageFont.truetype(p, size)
    return ImageFont.load_default()


def auto_vs_final(session, o):
    """Compare the untouched automatic mask with the current mask in IMAGE coordinates."""
    if not (o.get('auto_bbox') and o.get('bbox') and o.get('auto_mask_file')):
        return None
    a = session.load_mask(o, auto=True)
    f = session.load_mask(o)
    if a is None or f is None:
        return None
    ax0, ay0 = o['auto_bbox'][:2]
    fx0, fy0 = o['bbox'][:2]
    x0, y0 = min(ax0, fx0), min(ay0, fy0)
    x1 = max(ax0 + a.shape[1], fx0 + f.shape[1])
    y1 = max(ay0 + a.shape[0], fy0 + f.shape[0])
    A = np.zeros((y1 - y0, x1 - x0), bool)
    F = np.zeros_like(A)
    A[ay0 - y0:ay0 - y0 + a.shape[0], ax0 - x0:ax0 - x0 + a.shape[1]] = a
    F[fy0 - y0:fy0 - y0 + f.shape[0], fx0 - x0:fx0 - x0 + f.shape[1]] = f
    inter, union = int((A & F).sum()), int((A | F).sum())
    return dict(iou=round(inter / union, 4) if union else None, auto_px=int(A.sum()), final_px=int(F.sum()),
                area_change_pct=round(100.0 * (F.sum() - A.sum()) / max(A.sum(), 1), 2),
                added_by_reviewer_px=int((F & ~A).sum()), removed_by_reviewer_px=int((A & ~F).sum()))


def _overlap(Rm, rbox, m, mbox):
    """Intersection pixel count of a reference mask (box rbox = x0, y0) and an object mask."""
    rx0, ry0 = rbox
    mx0, my0 = mbox
    x0, y0 = max(rx0, mx0), max(ry0, my0)
    x1, y1 = min(rx0 + Rm.shape[1], mx0 + m.shape[1]), min(ry0 + Rm.shape[0], my0 + m.shape[0])
    if x1 <= x0 or y1 <= y0:
        return 0
    return int((Rm[y0 - ry0:y1 - ry0, x0 - rx0:x1 - rx0] & m[y0 - my0:y1 - my0, x0 - mx0:x1 - mx0]).sum())


IJ_COLS = (('Area', 'ImageJ Area'), ('Perim.', 'ImageJ Perim.'), ('Feret', 'ImageJ Feret'), ('Circ.', 'ImageJ Circ.'))
REF_HEAD = ['Original filename', 'Image ID', 'Reference ID', 'ROI name', 'ROI type', 'Reference file', 'Well / region ID',
            'Scale (µm per px, confirmed)',
            'Reference: pixel area (px²)', 'Reference: area, smoothed boundary (px²)', 'Reference: perimeter (px)',
            'Reference: ECD (px)', 'Reference: max Feret (px)', 'Reference: circularity',
            'Best automatic candidate', 'Its review status', 'IoU reference vs automatic mask',
            'Automatic candidates overlapping (n)',
            'Best accepted object', 'IoU reference vs final cell-covered', 'IoU reference vs final envelope',
            'Final cell-covered area vs reference (%)', 'Final envelope area vs reference (%)',
            'Final ECD vs reference (%)', 'Final max Feret vs reference (%)',
            'Final circularity (cell-covered) minus reference', 'Final circularity (envelope) minus reference',
            *[c[1] + ' (as in Results)' for c in IJ_COLS], 'Results row matched',
            'Same-ROI check: ImageJ Area / app pixel area of this ROI', 'Note']


def reference_rows(session, im, cal):
    """Per reference ROI: best automatic candidate and best accepted object, both by IoU; the
    reference measured with the app's own definitions; ImageJ's own values if a Results table
    was uploaded. Plus accepted objects that overlap no reference."""
    from . import references as REF
    refs = im.get('references') or []
    if not refs:
        return [], None
    W, H = im['width'], im['height']
    objs = [o for o in session.data['objects'].values() if o['image_id'] == im['image_id'] and o['bbox']]
    env_r = float(load_config()['measurement']['envelope_radius_px'])
    cache = {}

    def masks(o):
        if o['object_id'] not in cache:
            f = session.load_mask(o)
            a = session.load_mask(o, auto=True) if o.get('auto_mask_file') and o.get('auto_bbox') else None
            cache[o['object_id']] = (f, M.envelope_of(f, o.get('envelope_radius_px', env_r)) if f is not None else None, a)
        return cache[o['object_id']]

    files = {f['file_id']: f['original_filename'] for f in im.get('reference_files', [])}
    by_file = {}
    for r in refs:
        by_file.setdefault(r['file_id'], []).append(r)
    res = im.get('imagej_results')
    s = cal['um_per_px']
    rows, matched_acc, summ = [], set(), dict(n=len(refs), auto_found=0, final_matched=0, iou_auto=[], iou_cells=[],
                                               iou_env=[], results_matched=0)
    for r in refs:
        l, t, rr, b = r['bounds']
        bx0, by0 = max(0, int(np.floor(l)) - 1), max(0, int(np.floor(t)) - 1)
        bx1, by1 = min(W, int(np.ceil(rr)) + 1), min(H, int(np.ceil(b)) + 1)
        Rm = REF.reference_mask(r, (by1 - by0, bx1 - bx0), (bx0, by0))
        rpx = int(Rm.sum())
        note = []
        if r.get('partly_outside'):
            note.append('ROI partly outside the image')
        if not rpx:
            rows.append([im['original_filename'], im['image_id'], r['ref_id'], r['name'], r['kind'], files.get(r['file_id'], ''),
                         '', s] + [None] * (len(REF_HEAD) - 9) + ['ROI covers no pixel centre'])
            continue
        meas = M.measure_mask(np.pad(Rm, 2))
        cx, cy = (bx0 + np.nonzero(Rm)[1].mean() + 0.5), (by0 + np.nonzero(Rm)[0].mean() + 0.5)
        well = next((w['well_id'] for w in im['wells'] if _pip(cx, cy, w['polygon'])), '')
        best_a, best_f, n_auto = (None, 0.0), (None, 0.0, 0.0), 0
        for o in objs:
            ox0, oy0, ox1, oy1 = o['bbox']
            hit = lambda q: q[2] > bx0 and q[0] < bx1 and q[3] > by0 and q[1] < by1   # noqa: E731
            if not hit(o['bbox']) and not (o.get('auto_bbox') and hit(o['auto_bbox'])):
                continue
            f, env, a = masks(o)
            if o['source'] == 'auto' and a is not None:
                inter = _overlap(Rm, (bx0, by0), a, tuple(o['auto_bbox'][:2]))
                if inter:
                    n_auto += 1
                    iou = inter / (rpx + int(a.sum()) - inter)
                    if iou > best_a[1]:
                        best_a = (o, iou)
            if o['status'] == 'accepted' and f is not None:
                ic = _overlap(Rm, (bx0, by0), f, (ox0, oy0))
                ie = _overlap(Rm, (bx0, by0), env, (ox0, oy0))
                if ie:
                    iou_c = ic / (rpx + int(f.sum()) - ic)
                    iou_e = ie / (rpx + int(env.sum()) - ie)
                    if iou_e > best_f[2]:
                        best_f = (o, iou_c, iou_e)
        fo = best_f[0]
        if fo is not None and best_f[2] >= 0.1:
            matched_acc.add(fo['object_id'])
        ij, how = REF.results_row_for(r, by_file[r['file_id']], res)
        ij_vals = [ij.get(c[0]) if ij else None for c in IJ_COLS]
        if ij:
            summ['results_matched'] += 1
        ratio = round(ij['Area'] / rpx, 4) if ij and isinstance(ij.get('Area'), float) else None

        def pct(a, b):
            return round(100.0 * (a - b) / b, 2) if (a is not None and b) else None
        fm = (fo['measure_px'] if fo else None) or {}
        fc, fe = fm.get('cells') or {}, fm.get('envelope') or {}
        if best_a[1] >= 0.5:
            summ['auto_found'] += 1
        if fo is not None and max(best_f[1], best_f[2]) >= 0.5:
            summ['final_matched'] += 1
        if best_a[0] is not None:
            summ['iou_auto'].append(best_a[1])
        if fo is not None:
            summ['iou_cells'].append(best_f[1])
            summ['iou_env'].append(best_f[2])
        rows.append([im['original_filename'], im['image_id'], r['ref_id'], r['name'], r['kind'], files.get(r['file_id'], ''),
                     well, s, rpx, round(meas['area'], 2), round(meas['perimeter'], 2), round(meas['ecd'], 3),
                     round(meas['feret_max'], 3), round(meas['circularity'], 4),
                     best_a[0]['object_id'] if best_a[0] else '', best_a[0]['status'] if best_a[0] else '',
                     round(best_a[1], 4) if best_a[0] else None, n_auto,
                     fo['object_id'] if fo else '', round(best_f[1], 4) if fo else None, round(best_f[2], 4) if fo else None,
                     pct(fc.get('area'), meas['area']), pct(fe.get('area'), meas['area']),
                     pct(fc.get('ecd'), meas['ecd']), pct(fc.get('feret_max'), meas['feret_max']),
                     round(fc['circularity'] - meas['circularity'], 4) if fc else None,
                     round(fe['circularity'] - meas['circularity'], 4) if fe else None,
                     *ij_vals, how if res else '', ratio,
                     '; '.join(note + ([] if fo else ['no accepted object overlaps this reference']))])
    for o in objs:
        if o['status'] == 'accepted' and o['object_id'] not in matched_acc:
            rows.append([im['original_filename'], im['image_id'], '', '', '', '', o['well_id'] or '', s] +
                        [None] * 10 + [o['object_id']] + [None] * (len(REF_HEAD) - 20) +
                        ['accepted object that overlaps no reference outline (extra object, or not drawn in ImageJ)'])
            summ.setdefault('accepted_without_reference', 0)
            summ['accepted_without_reference'] += 1
    return rows, summ


def _pip(x, y, poly):
    inside = False
    for i in range(len(poly)):
        x1, y1 = poly[i - 1]
        x2, y2 = poly[i]
        if (y1 > y) != (y2 > y) and x < (x2 - x1) * (y - y1) / (y2 - y1) + x1:
            inside = not inside
    return inside


def annotate(session, image_id, out_path):
    """Full-resolution JPEG (quality 92): original + every object contour with its ID and status.
    A visual aid only; the data are the mask PNGs. (v0.5 wrote optimised PNGs: exporting a
    350-image session took 622 s and 469 MB; with JPEG and faster outlines 38 s and 131 MB.)"""
    im = session.image(image_id)
    rgb = load_rgb8(session.abspath(im['stored_path']))
    img = Image.fromarray(rgb).convert('RGB')
    d = ImageDraw.Draw(img)
    col = {'accepted': (0, 200, 0), 'rejected': (230, 0, 0), 'unreviewed': (255, 200, 0)}
    sigma = float(load_config()['measurement']['boundary_sigma_px'])
    f = _font(18)
    for w in im['wells']:
        pts = [tuple(p) for p in w['polygon']]
        d.line(pts + [pts[0]], fill=(80, 160, 255), width=2)
        d.text((pts[0][0] + 4, pts[0][1] + 2), w['well_id'], fill=(80, 160, 255), font=f)
    for o in sorted(session.data['objects'].values(), key=lambda o: o['object_id']):
        if o['image_id'] != image_id or not o['bbox']:
            continue
        m = session.load_mask(o)
        for c in contours_xy(m, (o['bbox'][0], o['bbox'][1]), sigma):
            pts = [tuple(p) for p in c]
            d.line(pts + [pts[0]], fill=col[o['status']], width=2 if o['status'] == 'accepted' else 1)
        x, y = o['bbox'][0], max(0, o['bbox'][1] - 20)
        label = f"{o['object_id']} ({o['status']})"
        d.text((x + 1, y + 1), label, fill=(0, 0, 0), font=f)
        d.text((x, y), label, fill=col[o['status']], font=f)
    img.save(out_path, quality=92)


def export_session(session):
    stamp = now().replace(':', '').replace('-', '')
    rel_dir = os.path.join('exports', stamp)
    out_dir = session.abspath(rel_dir)
    os.makedirs(os.path.join(out_dir, 'masks'), exist_ok=True)
    os.makedirs(os.path.join(out_dir, 'annotated'), exist_ok=True)
    cfg = load_config()
    data = session.data

    wb = Workbook()
    ws = wb.active
    ws.title = 'Measurements'
    head = ['Original filename', 'Image ID', 'Image SHA-256', 'Well / region ID', 'Region type', 'Object ID', 'Object source',
            'Review status', 'Centroid x (px)', 'Centroid y (px)', 'BBox x0 (px)', 'BBox y0 (px)',
            'BBox x1 (px, excl.)', 'BBox y1 (px, excl.)', 'Calibration ID', 'Calibration status',
            'Scale (µm per px)']
    for key, name in REGIONS:
        head += [f'{name}: area, smoothed boundary (px²)', f'{name}: raw pixel area (px²)', f'{name}: perimeter (px)',
                 f'{name}: ECD = 2·√(A/π) (px)', f'{name}: max Feret (px)', f'{name}: circularity 4πA/P²',
                 f'{name}: parts (n)', f'{name}: area, smoothed boundary (µm²)', f'{name}: raw pixel area (µm²)',
                 f'{name}: perimeter (µm)', f'{name}: ECD (µm)', f'{name}: max Feret (µm)']
    head += ['Cell-covered: pixels added by gap closing / hole filling (px²)',
             'Cell-covered: raw pixel area WITHOUT added pixels (px²; area only, no perimeter/circularity)',
             'Cell-covered: raw pixel area WITHOUT added pixels (µm²; area only)',
             'Parts: seed (n)', 'Parts: attached by proximity rule, unconfirmed (n)', 'Parts: assigned by reviewer (n)',
             'Approved by', 'Approved at', 'Approved mask version', 'Current mask version', 'Mask SHA-256 (16)',
             'Auto BBox x0 (px)', 'Auto BBox y0 (px)', 'Auto BBox x1 (px, excl.)', 'Auto BBox y1 (px, excl.)',
             'IoU auto vs final mask', 'Raw area change vs auto (%)',
             'Envelope closing radius (px)', 'Boundary smoothing sigma (px)',
             'Primary region', 'Quality flags', 'Notes', 'Mask file (crop, see BBox)', 'Auto-prediction mask file',
             'Annotated image file', 'Profile', 'Segmenter', 'Settings hash', 'Well source', 'Well detection',
             'Image used for tuning']
    ws.append(head)
    log = []
    images_with_rows = set()
    n_rows = 0
    for im in data['images']:
        iid = im['image_id']
        cal = session.calibration_state(iid)
        calrec = data['calibrations'].get(im['calibration_id']) if im['calibration_id'] else None
        objs = sorted([o for o in data['objects'].values() if o['image_id'] == iid], key=lambda o: o['object_id'])
        ann_rel = os.path.join('annotated', f"{iid}__{os.path.splitext(disk_name(im['original_filename']))[0]}.jpg")
        well_src = {w['well_id']: w['source'] for w in im['wells']}
        kind = {w['well_id']: REGION_TXT[region_kind(w)] for w in im['wells']}
        wd = im.get('well_detection') or {}
        wd_txt = f"{wd.get('method', '')} {wd.get('status', '')}".strip() if wd else ('example preset' if im.get('preset_applied') else 'manual')
        if wd.get('status') == 'failed':
            log.append([im['original_filename'], iid, '', '', 'wells', 'automatic well detection failed', '',
                        wd.get('reason', ''), ''])
        for w in im['wells']:
            det = w.get('detection') or {}
            if det.get('grid_cell') and w['source'] == 'auto (grid cell)':
                log.append([im['original_filename'], iid, w['well_id'], '', 'grid cell',
                            'grid cell proposed (no well outline found)', '',
                            'working region, not a well boundary; it may hold several objects or none, and objects '
                            'can be cut by its edge', ''])
            elif det and w['source'] == 'auto' and (det.get('partial') or (det.get('rel_score') or 1) < 0.35):
                log.append([im['original_filename'], iid, w['well_id'], '', 'well',
                            'auto well unedited: ' + ('cut by image edge' if det.get('partial') else 'low confidence'),
                            '', f"relative outline score {det.get('rel_score')}", ''])
        # --- image level log entries
        if im['processing'].get('status') != 'processed':
            log.append([im['original_filename'], iid, '', '', 'image', im['processing'].get('status', 'not_processed'),
                        '', im['processing'].get('message', '') or im['processing'].get('error', ''), ''])
        if not im['review_done']:
            log.append([im['original_filename'], iid, '', '', 'image', 'image review not marked complete', '', '', ''])
        for f in im['image_flags']:
            log.append([im['original_filename'], iid, '', '', 'image flag', f['flag'], f.get('set_by', ''),
                        f.get('reason', ''), ''])
        if cal['status'] != 'confirmed':
            log.append([im['original_filename'], iid, '', '', 'calibration', f"calibration {cal['status']}",
                        '', 'physical measurements not reported for this image', ''])
        for w in im['wells']:
            has = [o for o in objs if o['well_id'] == w['well_id']]
            if w['no_measurable_object']:
                log.append([im['original_filename'], iid, w['well_id'], '', 'well', 'marked: no measurable object',
                            '', w.get('note', ''), ''])
            elif not has and im['processing'].get('status') == 'processed':
                log.append([im['original_filename'], iid, w['well_id'], '', 'well', 'no candidate found', '',
                            w.get('note', ''), ''])
        for o in objs:
            flags = [k for k, v in session.object_view(o, cal)['flags'].items() if v]
            if o['status'] != 'accepted':
                log.append([im['original_filename'], iid, o['well_id'] or '', o['object_id'], 'object', o['status'],
                            ', '.join(flags), o['notes'], o['mask_file']])
                continue
            # --- accepted object row
            images_with_rows.add(iid)
            mask_rel = os.path.join('masks', f"{o['object_id']}.png")
            shutil.copyfile(session.abspath(o['mask_file']), os.path.join(out_dir, mask_rel))
            auto_rel = ''
            if o.get('auto_mask_file') and os.path.exists(session.abspath(o['auto_mask_file'])):
                auto_rel = os.path.join('masks', f"{o['object_id']}_auto.png")
                shutil.copyfile(session.abspath(o['auto_mask_file']), os.path.join(out_dir, auto_rel))
            src = 'manual' if o['source'] == 'manual' else ('auto + edited' if o['edited'] else 'auto')
            row = [im['original_filename'], iid, im['sha256'], o['well_id'] or '', kind.get(o['well_id'], ''),
                   o['object_id'], src, o['status'],
                   o['centroid'][0], o['centroid'][1], *o['bbox'], im['calibration_id'] or '',
                   cal['status'] if cal['status'] == 'confirmed' else f"{cal['status']} — physical values not reported",
                   cal['um_per_px'] if cal['um_per_px'] else None]
            s2 = cal['um_per_px'] ** 2 if cal['um_per_px'] else None
            for key, _ in REGIONS:
                px = o['measure_px'][key]
                um = M.to_physical(px, cal['um_per_px'])
                row += [px['area'], px['mask_pixels'], px['perimeter'], px['ecd'], px['feret_max'], px['circularity'],
                        px['n_parts']]
                row += ([um['area'], px['mask_pixels'] * s2, um['perimeter'], um['ecd'], um['feret_max']]
                        if um else [None] * 5)
            raw = o['measure_px']['cells']['mask_pixels']
            added = int(o.get('added_px') or 0)
            hows = [p.get('how') for p in o.get('parts', [])]
            ap = o.get('approval') or {}
            av = auto_vs_final(session, o)
            row += [added, raw - added, (raw - added) * s2 if s2 else None,
                    hows.count('seed'), hows.count('proximity'), hows.count('user'),
                    ap.get('by') or '(not recorded)', ap.get('at', ''), ap.get('mask_version'), o.get('mask_version'),
                    o.get('mask_sha256') or '', *(o.get('auto_bbox') or [None] * 4),
                    av['iou'] if av else None, av['area_change_pct'] if av else None,
                    o.get('envelope_radius_px'), o.get('boundary_sigma_px'),
                    'Cell-covered (see Settings & definitions)', ', '.join(flags), o['notes'], mask_rel, auto_rel,
                    ann_rel, o['settings'].get('profile') or '', o['settings'].get('segmenter') or '',
                    o['settings'].get('params_hash') or '', well_src.get(o['well_id'], ''),
                    wd_txt, 'yes' if im['used_for_tuning'] else 'no']
            ws.append(row)
            n_rows += 1
        if objs:
            annotate(session, iid, os.path.join(out_dir, ann_rel))
    _style(ws, widths={1: 34, 2: 10, 3: 18}, wrap_header=True)
    ws.freeze_panes = 'F2'
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if isinstance(c.value, float):
                c.number_format = '0.000' if 'circularity' in str(ws.cell(1, c.column).value) else '#,##0.00'

    wl = wb.create_sheet('Review log')
    wl.append(['Original filename', 'Image ID', 'Well / region ID', 'Object ID', 'Item', 'Status / issue', 'Flags / set by',
               'Notes / reason', 'Mask file (session folder)'])
    for r in log:
        wl.append(r)
    _style(wl, widths={1: 34, 6: 34, 7: 40, 8: 70, 9: 36}, wrap_header=False)

    wc = wb.create_sheet('Calibrations')
    wc.append(['Calibration ID', 'Status', 'µm per px', 'Known length', 'Unit', 'Line length (px)',
               'Line x1', 'Line y1', 'Line x2', 'Line y2', 'Drawn on image', 'Applied to images',
               'Compatibility confirmed for', 'Created', 'Confirmed at', 'Note'])
    for c in data['calibrations'].values():
        (x1, y1), (x2, y2) = c['line']
        wc.append([c['cal_id'], c['status'], c['um_per_px'], c['known_length'], c['unit'], c['line_length_px'],
                   x1, y1, x2, y2, c['source_image_id'], ', '.join(c['applied_to']),
                   ', '.join(c['compatibility_confirmed_for']), c['created'], c['confirmed_at'] or '', c['note']])
    _style(wc, widths={16: 80})

    wa = wb.create_sheet('Auto vs final')
    wa.append(['Original filename', 'Image ID', 'Well / region ID', 'Object ID', 'Source', 'Review status', 'Mask version',
               'Mask edits (n)', 'IoU auto vs final', 'Auto raw area (px²)', 'Final raw area (px²)',
               'Area change (%)', 'Pixels added by reviewer', 'Pixels removed by reviewer',
               'Proximity parts still attached (n)', 'Auto BBox', 'Final BBox', 'Note'])
    for im in data['images']:
        for o in sorted([o for o in data['objects'].values() if o['image_id'] == im['image_id']],
                        key=lambda o: o['object_id']):
            note = AV_NOTE.get((o['source'] == 'auto', o['status']), '')
            # correction numbers only for ACCEPTED automatic objects: a wrong candidate rejected without
            # editing would otherwise show IoU 1 (its mask was never changed)
            av = auto_vs_final(session, o) if (o['source'] == 'auto' and o['status'] == 'accepted') else None
            edits = sum(1 for h in o.get('history', []) if h['action'].split(' ')[0].split('_')[0] in
                        ('brush', 'polygon', 'assign', 'unassign'))
            wa.append([im['original_filename'], im['image_id'], o['well_id'] or '', o['object_id'], o['source'],
                       o['status'], o.get('mask_version'), edits, av['iou'] if av else None,
                       av['auto_px'] if av else None, av['final_px'] if av else (o['measure_px'] or {}).get('cells', {}).get('mask_pixels'),
                       av['area_change_pct'] if av else None, av['added_by_reviewer_px'] if av else None,
                       av['removed_by_reviewer_px'] if av else None,
                       sum(1 for p in o.get('parts', []) if p.get('how') == 'proximity'),
                       json.dumps(o.get('auto_bbox')), json.dumps(o.get('bbox')), note])
    _style(wa, widths={1: 34, 16: 22, 17: 22, 18: 70})

    we = wb.create_sheet('Review effort')
    we.append(EFFORT_HEAD)
    for im in data['images']:
        ef = session.review_effort(im['image_id'])
        objs = [o for o in data['objects'].values() if o['image_id'] == im['image_id']]
        auto = [o for o in objs if o['source'] == 'auto']
        acc = [o for o in auto if o['status'] == 'accepted']
        ious = [a['iou'] for a in (auto_vs_final(session, o) for o in acc) if a and a['iou'] is not None]
        manual = [o for o in objs if o['source'] == 'manual']
        from_part = [o for o in manual if str((o.get('settings') or {}).get('segmenter', '')).startswith('part')]
        we.append([im['original_filename'], im['image_id'], 'yes' if im['review_done'] else 'no',
                   ef['active_review_s'], ef['active_window_s'], ef['active_seconds_est'], ef['idle_cap_s'],
                   ef['events'], ef['first'] or '', ef['last'] or '', len(auto),
                   sum(1 for o in acc if not o['edited']), sum(1 for o in acc if o['edited']),
                   sum(1 for o in auto if o['status'] == 'rejected'),
                   sum(1 for o in objs if o['status'] == 'unreviewed'),
                   len(manual) - len(from_part), len(from_part),
                   sum(1 for o in manual if o['status'] == 'accepted'),
                   round(float(np.mean(ious)), 4) if ious else None, sum(1 for v in ious if v < 0.9),
                   json.dumps(ef['counts'])])
    _style(we, widths={1: 34, 4: 16, 6: 16, 21: 60})

    if any(im.get('references') for im in data['images']):
        wr = wb.create_sheet('Reference comparison')
        wr.append(REF_HEAD)
        ws_sum = wb.create_sheet('Reference summary')
        ws_sum.append(['Original filename', 'Image ID', 'Reference outlines (n)', 'ROIs skipped (n)',
                       'Found automatically (IoU ≥ 0.5)', 'Matched by an accepted object (IoU ≥ 0.5)',
                       'Mean best IoU, automatic', 'Mean IoU, final cell-covered', 'Mean IoU, final envelope',
                       'Accepted objects without reference (n)', 'Results rows matched', 'Reference files'])
        for im in data['images']:
            rows, summ = reference_rows(session, im, session.calibration_state(im['image_id']))
            for r in rows:
                wr.append(r)
            if summ:
                mean = lambda v: round(float(np.mean(v)), 4) if v else None   # noqa: E731
                ws_sum.append([im['original_filename'], im['image_id'], summ['n'],
                               sum(len(f.get('skipped') or []) for f in im.get('reference_files', [])),
                               summ['auto_found'], summ['final_matched'], mean(summ['iou_auto']),
                               mean(summ['iou_cells']), mean(summ['iou_env']), summ.get('accepted_without_reference', 0),
                               summ['results_matched'],
                               '; '.join(f"{f['original_filename']} (sha256 {f['sha256'][:12]})" for f in im['reference_files'])])
        _style(wr, widths={1: 34, 41: 60})
        _style(ws_sum, widths={1: 34, 12: 60})

    wp = wb.create_sheet('Processing parameters')
    wp.append(['Original filename', 'Image ID', 'Profile', 'Segmenter', 'Settings hash', 'Detection time',
               'Segmenter + grouping parameters (JSON)', 'Per-well values (well size, grouping in px, background levels)',
               'Well detection method', 'Well detection status', 'Well detection parameters (JSON)', 'Well detection info'])
    for im in data['images']:
        pr = im.get('processing') or {}
        wd = im.get('well_detection') or {}
        wp.append([im['original_filename'], im['image_id'], pr.get('profile', ''), pr.get('segmenter', ''),
                   pr.get('params_hash', ''), pr.get('time', ''), json.dumps(pr.get('params', {}), ensure_ascii=False),
                   json.dumps(pr.get('per_well', {}), ensure_ascii=False)[:32000], wd.get('method', ''),
                   wd.get('status', 'not run' if not im.get('preset_applied') else 'example preset'),
                   json.dumps(wd.get('params', {}), ensure_ascii=False), json.dumps(wd.get('info', {}), ensure_ascii=False)])
    _style(wp, widths={1: 34, 7: 60, 8: 60, 11: 60, 12: 50})

    wsd = wb.create_sheet('Settings & definitions')
    rows = [('Export', f"{now()} — app {APP_VERSION} — session '{data['name']}'"),
            ('Accepted objects exported', n_rows),
            ('Coordinates', 'Continuous image pixel coordinates of the ORIGINAL image: x right, y down; pixel (row r, col c) '
                            'covers [c, c+1) x [r, r+1). Display zoom never affects stored coordinates or measurements.'),
            ('Primary region', 'Cell-covered region. Envelope and dense core are reported alongside, each with its own '
                               'area AND perimeter (circularity always uses area and perimeter of the same region).'),
            ('Cell-covered region', M.REGION_DEFINITIONS['cells']),
            ('Envelope region', M.REGION_DEFINITIONS['envelope']),
            ('Dense core region', M.REGION_DEFINITIONS['core']),
            ('Boundary', f"All values come from one sub-pixel boundary: Gaussian blur (sigma = "
                         f"{cfg['measurement']['boundary_sigma_px']} px) of the binary mask, 0.5 iso-contour. "
                         'Holes inside a part are excluded from area; their outlines count in perimeter. '
                         'Multi-part regions: areas and perimeters summed, Feret across all parts.'),
            ('Area', 'Two values per region: "area, smoothed boundary" = polygon area of the sub-pixel boundary used '
                     'for perimeter/circularity (outer minus holes); "raw pixel area" = number of mask pixels. '
                     'They differ by a fraction of a pixel per boundary pixel.'),
            ('Added pixels', 'Pixels inside the cell-covered region that did NOT pass the cell threshold but were added '
                             'by gap closing (2 px) or hole filling. Reported separately; the raw area without them is '
                             'also given. That area is an AREA ONLY: perimeter, ECD, Feret and circularity always '
                             'describe the cell-covered region including added pixels. Do not divide the area without '
                             'added pixels by that perimeter: the two belong to different regions, so the resulting '
                             '"circularity" would be meaningless.'),
            ('Removed parts', 'Pixels that passed the cell threshold but were removed by a clean-up step stay available '
                              'as parts with a reason: wall_shape (wall/shadow-shaped component), wall_band (inside the '
                              'wall margin), thin (removed by opening), small (below the speck limit), split_line (the '
                              'cut between two pieces of a component split by the optional watershed profile; it also '
                              'holds any added pixels on the cut, so assigning it back restores the unsplit mask '
                              'exactly). None is part of an object unless the reviewer assigns it.'),
            ('Proximity parts', 'Small parts attached to a candidate only because they lay within the proximity '
                                'distance. Unconfirmed suggestion until the reviewer accepts the object (the count at '
                                'approval is stored).'),
            ('Approval', 'Approved by / at / mask version: who accepted the object, when, and which mask version. '
                         'Any later mask edit resets the object to unreviewed.'),
            ('Auto vs final', 'IoU and area change between the untouched automatic mask (Auto BBox) and the final mask, '
                              'in image coordinates, for ACCEPTED automatic objects only. It measures how much the '
                              'reviewer corrected an accepted candidate, NOT accuracy. Rejected candidates get no IoU '
                              '(an unchanged rejected mask would give IoU 1). Judge the automatic step from all of: '
                              'accepted unchanged / accepted after edit / rejected candidates, objects the reviewer '
                              'added (missed), and these IoU values (Review effort sheet, per image).'),
            ('Review effort', 'Two time figures per image. "Measured active review time": counted by the app while the '
                              'image was open in a visible, focused window and the reviewer moved, clicked, typed or '
                              'scrolled within the previous 30 s (empty if no browser reported time). "Rough estimate '
                              'from action gaps": sum of gaps between consecutive reviewer actions, each capped at 120 s; '
                              'it misses thinking time before the first action and can include time spent elsewhere. '
                              'Neither measures attention exactly.'),
            ('Reference comparison', 'Outlines drawn independently in ImageJ and uploaded as references (RoiSet .zip or '
                                     '.roi; optional Results .csv). They are never objects. Each reference is rasterised '
                                     'by pixel centre and measured with THIS app\'s definitions (one sub-pixel boundary, '
                                     'sigma 1 px), so differences to the final object are differences of outline, not of '
                                     'formula. "Best automatic candidate" = the automatic mask with the highest IoU '
                                     '(segmentation accuracy against this reference); "best accepted object" = the accepted '
                                     'object whose envelope has the highest IoU. Compare with the cell-covered region when '
                                     'the reference traces cells only, with the envelope when it encloses the whole cluster. '
                                     'ImageJ\'s own values are copied from the Results table as they are (its units, its '
                                     'perimeter definition: polygon length, no smoothing; ImageJ caps circularity at 1). '
                                     'Same-ROI check: ImageJ Area / pixel count of the same ROI here; 1.0 means the ROI was '
                                     'read and rasterised as ImageJ did (in pixel units).'),
            ('Perimeter', 'Total boundary length (outer + hole outlines, all parts).'),
            ('ECD', 'Equivalent circular diameter = 2 * sqrt(area / pi).'),
            ('Max Feret', 'Largest distance between two boundary points (across all parts of the region).'),
            ('Circularity', '4 * pi * area / perimeter^2 of the SAME region. Not clipped. For multi-part regions it '
                            'is not a single-object shape descriptor (see parts count).'),
            ('Calibration', 'Physical columns are filled only when the image calibration is CONFIRMED (Set Scale line '
                            'drawn/reviewed by the user and, for other images, compatibility confirmed). Unconfirmed '
                            'estimates are never used.'),
            ('2D projection', 'All values describe the 2D image projection of the object.'),
            ('Validation status', 'Measurement code verified on synthetic shapes (tests/test_measure.py). Segmentation '
                                  'accuracy is NOT validated; needs researcher-approved masks on images not used for tuning.'),
            ('Well / region ID', 'W1..Wn = wells: outline found by classical detection (app/wells.py), example preset, '
                                 'copied, or drawn by hand. G1..Gn = grid cells: when no well outline is found but dark '
                                 'objects sit on a regular grid, square WORKING REGIONS are proposed one per grid period. '
                                 'A grid cell is not a well boundary and not a count: it may contain several objects or '
                                 'none, and an object can be cut by its edge. See the "Region type" and "Well source" '
                                 'columns; parameters are in the Processing parameters sheet.'),
            ('Measurement settings', json.dumps(cfg['measurement'], ensure_ascii=False)),
            ]
    for name, prof in cfg['profiles'].items():
        rows.append((f'Profile {name}', json.dumps(prof, ensure_ascii=False)))
    wsd.append(['Item', 'Value'])
    for r in rows:
        wsd.append(list(r))
    _style(wsd, widths={1: 28, 2: 140})
    for row in wsd.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical='top')

    xlsx_rel = os.path.join(rel_dir, 'spheroid_measurements.xlsx')
    wb.save(session.abspath(xlsx_rel))
    with open(os.path.join(out_dir, 'session_snapshot.json'), 'w', encoding='utf-8') as f:
        f.write(json.dumps(data, ensure_ascii=False, indent=1))
    shutil.copyfile(os.path.join(os.path.dirname(__file__), '..', 'config', 'profiles.json'),
                    os.path.join(out_dir, 'profiles_snapshot.json'))
    return dict(folder=rel_dir, xlsx=xlsx_rel, accepted_rows=n_rows, review_log_rows=len(log))


def _style(ws, widths=None, wrap_header=True):
    for c in ws[1]:
        c.font = BOLD
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=wrap_header, vertical='top')
    if wrap_header:
        ws.row_dimensions[1].height = 60
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = FONT
    for col in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col)].width = (widths or {}).get(col, 14)
